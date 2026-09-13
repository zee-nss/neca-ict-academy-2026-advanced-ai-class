import pandas as pd
import matplotlib.pyplot as plt
from datetime import datetime
import numpy as np
import os


# --------------------------------------------------------
# SETUP
# --------------------------------------------------------
plt.style.use('seaborn-v0_8-whitegrid')

COLORS = ['#2C3E50', '#E74C3C', '#3498DB', '#2ECC71', '#F39C12']



# Load our cleaned data
base = "C:/Users/zeena/Desktop/neca ict academy 2026 advanced ai class"
df = pd.read_csv(os.path.join(base, "sales_data_CLEANED.csv"))
df['date'] = pd.to_datetime(df['date'], dayfirst=False)

print("Loaded", len(df), "rows of sales data")
print("Date range:", df['date'].min().date(), "to", df['date'].max().date())



# --------------------------------------------------------
# PART 1: CREATE ALL REPORT CHARTS
# --------------------------------------------------------

def create_report_charts(df):
    """
    Generate all 6 charts for the weekly sales report.
    Returns the file path to the saved chart image.
    """
    
    # Create a 2x3 grid of subplots
    fig, axes = plt.subplots(2, 3, figsize=(16, 10))
    
    # Main title
    fig.suptitle('WEEKLY SALES PERFORMANCE REPORT', 
                 fontsize=18, fontweight='bold', y=0.98)

    # --------------------------------------------------
    # Chart 1: Revenue by Region (Pie Chart)
    # --------------------------------------------------
    region_data = df.groupby('region')['total_sale'].sum()
    
    axes[0,0].pie(
        region_data.values, 
        labels=region_data.index, 
        autopct='%1.1f%%',
        colors=COLORS, 
        startangle=90
    )
    axes[0,0].set_title('Revenue Distribution by Region', fontweight='bold')

    # --------------------------------------------------
    # Chart 2: Top 5 Products (Horizontal Bar)
    # --------------------------------------------------
    product_data = df.groupby('product')['total_sale'].sum()
    product_data = product_data.sort_values().tail(5)
    
    axes[0,1].barh(product_data.index, product_data.values, color=COLORS[2])
    axes[0,1].set_title('Top 5 Products by Revenue', fontweight='bold')
    axes[0,1].set_xlabel('Revenue ($)')

    # --------------------------------------------------
    # Chart 3: Daily Revenue Trend (Line)
    # --------------------------------------------------
    daily = df.groupby('date')['total_sale'].sum()
    
    axes[0,2].plot(
        daily.index, 
        daily.values, 
        color=COLORS[0], 
        linewidth=2, 
        marker='o',
        markersize=4
    )
    axes[0,2].set_title('Daily Revenue Trend', fontweight='bold')
    axes[0,2].set_ylabel('Revenue ($)')
    axes[0,2].tick_params(axis='x', rotation=45)

    # --------------------------------------------------
    # Chart 4: Sales by Day of Week (Bar)
    # --------------------------------------------------
    dow_order = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 
                 'Friday', 'Saturday', 'Sunday']
    dow_data = df.groupby('day_of_week')['total_sale'].sum()
    dow_data = dow_data.reindex([d for d in dow_order if d in dow_data.index])
    
    axes[1,0].bar(dow_data.index, dow_data.values, color=COLORS[3])
    axes[1,0].set_title('Revenue by Day of Week', fontweight='bold')
    axes[1,0].tick_params(axis='x', rotation=45)
    axes[1,0].set_ylabel('Revenue ($)')

    # --------------------------------------------------
    # Chart 5: Order Size vs Revenue (Scatter)
    # --------------------------------------------------
    axes[1,1].scatter(
        df['quantity'], 
        df['total_sale'], 
        alpha=0.5,
        c=COLORS[4], 
        edgecolors='white'
    )
    axes[1,1].set_title('Order Size vs Revenue', fontweight='bold')
    axes[1,1].set_xlabel('Quantity')
    axes[1,1].set_ylabel('Revenue ($)')

    # --------------------------------------------------
    # Chart 6: Key Metrics Summary (Text Panel)
    # --------------------------------------------------
    axes[1,2].axis('off')
    
    total_revenue = df['total_sale'].sum()
    total_orders = len(df)
    avg_order = df['total_sale'].mean()
    top_region = region_data.idxmax()
    top_product = product_data.idxmax()
    start_date = df['date'].min().date()
    end_date = df['date'].max().date()
    
    metrics_text = (
        "KEY METRICS\n"
        "\n"
        "Total Revenue: $" + format(int(total_revenue), ',') + "\n"
        "Total Orders: " + format(total_orders, ',') + "\n"
        "Avg Order Value: $" + format(int(avg_order), ',') + "\n"
        "\n"
        "Top Region: " + str(top_region) + "\n"
        "Top Product: " + str(top_product) + "\n"
        "\n"
        "Period: " + str(start_date) + "\n"
        "       to " + str(end_date)
    )
    
    axes[1,2].text(
        0.1, 0.9, 
        metrics_text, 
        transform=axes[1,2].transAxes,
        fontsize=12, 
        verticalalignment='top', 
        fontfamily='monospace',
        bbox=dict(boxstyle='round', facecolor='wheat', alpha=0.5)
    )

    # --------------------------------------------------
    # Save the chart panel
    # --------------------------------------------------
    plt.tight_layout()
    
    chart_filename = 'report_charts_' + datetime.now().strftime("%Y%m%d") + '.png'
    chart_path = os.path.join(base, chart_filename)
    plt.savefig(chart_path, dpi=150, bbox_inches='tight', facecolor='white')
    print("Charts saved:", chart_path)
    
    return chart_path


# --------------------------------------------------------
# PART 2: EXCEL REPORT GENERATOR
# --------------------------------------------------------

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage


def create_excel_report(df, chart_path):
    """
    Create a multi-sheet, professionally formatted Excel report.
    Returns the file path to the saved report.
    """
    
    wb = Workbook()

    # --------------------------------------------------
    # SHEET 1: EXECUTIVE SUMMARY
    # --------------------------------------------------
    ws = wb.active
    ws.title = "Executive Summary"
    
    # Title row
    ws.merge_cells('A1:F1')
    ws['A1'] = 'WEEKLY SALES PERFORMANCE REPORT'
    ws['A1'].font = Font(size=18, bold=True, color='FF1B4F72')
    ws['A1'].alignment = Alignment(horizontal='center')

    # Report metadata
    ws['A3'] = 'Report Generated:'
    ws['B3'] = datetime.now().strftime('%B %d, %Y at %H:%M')
    ws['A4'] = 'Period:'
    ws['B4'] = str(df['date'].min().date()) + ' to ' + str(df['date'].max().date())
    ws['A5'] = 'Data Rows Processed:'
    ws['B5'] = format(len(df), ',')

    # Key metrics section header
    ws['A7'] = 'KEY METRICS'
    ws['A7'].font = Font(size=14, bold=True, color='FF1B4F72')
    
    # Metrics with proper number formatting
    metrics = [
        ('Total Revenue', df['total_sale'].sum(), '$#,##0'),
        ('Total Orders', len(df), '#,##0'),
        ('Average Order Value', df['total_sale'].mean(), '$#,##0'),
        ('Maximum Single Order', df['total_sale'].max(), '$#,##0'),
        ('Unique Products Sold', df['product'].nunique(), '#,##0'),
        ('Active Regions', df['region'].nunique(), '#,##0'),
    ]
    
    for i, (label, value, fmt) in enumerate(metrics, start=8):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=value)
        cell.number_format = fmt

    # Insert the chart image
    img = XLImage(chart_path)
    img.width = 600
    img.height = 400
    ws.add_image(img, 'D7')

    # --------------------------------------------------
    # SHEET 2: REGIONAL BREAKDOWN
    # --------------------------------------------------
    ws2 = wb.create_sheet("Regional Breakdown")
    
    # Create the summary data
    region_summary = df.groupby('region').agg({
        'total_sale': ['sum', 'mean', 'count'],
        'quantity': 'sum'
    }).round(2)
    region_summary.columns = ['Total Revenue', 'Avg Order', 'Order Count', 'Total Quantity']
    
    # Sheet title
    ws2['A1'] = 'REGIONAL BREAKDOWN'
    ws2['A1'].font = Font(size=14, bold=True, color='FF1B4F72')
    
    # Styled header row
    headers = ['Region'] + list(region_summary.columns)
    for j, col in enumerate(headers, start=1):
        cell = ws2.cell(row=3, column=j, value=col)
        cell.font = Font(bold=True, color='FFFFFFFF')
        cell.fill = PatternFill(start_color='FF1B4F72', end_color='FF1B4F72', fill_type='solid')
    
    # Data rows
    for i, (region, row) in enumerate(region_summary.iterrows(), start=4):
        ws2.cell(row=i, column=1, value=region)
        for j, value in enumerate(row, start=2):
            ws2.cell(row=i, column=j, value=value)

    # --------------------------------------------------
    # SHEET 3: CLEANED DATA
    # --------------------------------------------------
    ws3 = wb.create_sheet("Cleaned Data")
    
    # Header row
    for j, col in enumerate(df.columns, start=1):
        cell = ws3.cell(row=1, column=j, value=col)
        cell.font = Font(bold=True)
    
    # All data rows
    for i, row in enumerate(df.values, start=2):
        for j, value in enumerate(row, start=1):
            ws3.cell(row=i, column=j, value=value)

    # Adjust column widths
    for col in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[col].width = 22

    # Save the workbook
    report_filename = 'weekly_sales_report_' + datetime.now().strftime("%Y%m%d") + '.xlsx'
    report_path = os.path.join(base, report_filename)
    wb.save(report_path)
    print("Excel report saved:", report_path)

    return report_path

# --------------------------------------------------------
# RUN THE REPORT GENERATOR
# --------------------------------------------------------
print()
print("Generating charts...")
chart_path = create_report_charts(df)

print("Creating Excel report...")
report_path = create_excel_report(df, chart_path)

print()
print("Report generation complete!")
print("Chart file:", chart_path)
print("Excel file:", report_path)

